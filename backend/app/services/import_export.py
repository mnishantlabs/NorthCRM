from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal

from app.models.business import BusinessStatus

_FIELD_ALIASES = {
    "business name": "business_name",
    "businessname": "business_name",
    "owner name": "owner_name",
    "ownername": "owner_name",
    "owner": "owner_name",
    "google maps url": "google_maps_url",
    "google maps link": "google_maps_url",
    "google maps": "google_maps_url",
    "assigned agent": "assigned_agent_id",
    "agent": "assigned_agent_id",
    "agent id": "assigned_agent_id",
    "agent id (uuid)": "assigned_agent_id",
    "status": "status",
}


def _normalize_key(key: str) -> str:
    if key is None:
        return ""
    normalized = str(key).strip().lower().replace("-", " ").replace("_", " ")
    normalized = re.sub(r"\s+", " ", normalized)
    return _FIELD_ALIASES.get(normalized, normalized.replace(" ", "_"))


def _clean(value):
    if value is None:
        return None
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return value
    text = str(value).strip()
    return text or None


def parse_csv(file_bytes: bytes) -> list[dict]:
    reader = csv.DictReader(io.StringIO(file_bytes.decode("utf-8-sig")))
    rows = []
    for raw in reader:
        if raw is None:
            continue
        rows.append({_normalize_key(k): _clean(v) for k, v in raw.items()})
    return rows


def parse_excel(file_bytes: bytes) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(min_row=1, max_row=1, values_only=True)
    try:
        first_row = next(iterator)
    except StopIteration:
        return []
    header = [_normalize_key(cell) for cell in first_row]

    rows = []
    for line in sheet.iter_rows(min_row=2, values_only=True):
        row = {}
        for index, key in enumerate(header):
            if not key:
                continue
            if index < len(line):
                row[key] = _clean(line[index])
        rows.append(row)
    workbook.close()
    return rows


def validate_rows(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    valid: list[dict] = []
    errors: list[dict] = []

    for index, row in enumerate(rows, start=2):
        problems: list[str] = []
        cleaned: dict = {key: value for key, value in row.items() if value is not None}

        name = str(cleaned.get("business_name") or "").strip()
        if not name:
            problems.append("business_name is required and cannot be empty")

        email = str(cleaned.get("email") or "").strip()
        if email and not (email.count("@") == 1 and email.split("@")[1].count(".") >= 1):
            problems.append(f"invalid email '{email}'")

        status_value = cleaned.get("status")
        if status_value is not None:
            try:
                cleaned["status"] = BusinessStatus(str(status_value).lower())
            except ValueError:
                problems.append(f"invalid status '{status_value}'")

        if problems:
            errors.append({"row": index, "errors": problems})
            continue

        cleaned["business_name"] = name
        if email:
            cleaned["email"] = email.lower()
        valid.append(cleaned)

    return valid, errors


def _stringify(value):
    if value is None:
        return ""
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, int | float | Decimal):
        return value
    return str(value)


def export_csv(rows: list[dict]) -> bytes:
    if not rows:
        return b""
    fieldnames = list(rows[0].keys())
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: _stringify(value) for key, value in row.items()})
    return output.getvalue().encode("utf-8-sig")


def export_excel(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    if rows:
        fieldnames = list(rows[0].keys())
        sheet.append(fieldnames)
        for row in rows:
            sheet.append([_stringify(row.get(field)) for field in fieldnames])
        for index in range(1, len(fieldnames) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
